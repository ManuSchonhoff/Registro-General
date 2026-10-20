"""Lectura del libro "Registro - USDT General" (Excel).

Arma filas planas de las hojas mensuales (ventas y compras), los bloques de RESULTADO USDT GENERAL y la ganancia de
cada venta de CLIENTES MARIANO. Las reglas están en `analisis/SPEC.md` §3. No modifica el archivo original.
"""
import calendar
import datetime as dt
import re
import statistics
import unicodedata
import warnings
from collections import Counter, defaultdict

import openpyxl

warnings.filterwarnings("ignore")

MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "oct": 10,
         "nov": 11, "dic": 12}
# Columnas (base 1) de cada lado: fecha, id, bruto, comisión, ars, canal, banco
LADOS = {"V": (1, 2, 3, 4, 7, 8, 9), "C": (11, 12, 13, 14, 17, 18, 19)}
VENTANA_DIAS = 7          # una fecha a más de 7 días del mes de la hoja se considera mal tipeada
PRECIO_ID = (900, 2500)   # cociente pesos / número del ID que indica que el ID es en realidad el USDT

CANALES = {
    "binance": "Binance", "fiwind": "Fiwind", "fwind": "Fiwind", "wallet cg": "Wallet Cambio General",
    "transf (cliente)": "Transferencia de cliente", "bybit": "Bybit", "astropay": "AstroPay",
}
BANCOS = {
    "personal pay": "Personal Pay", "brubank": "Brubank", "bna+": "Banco Nación (BNA+)", "uala": "Ualá",
    "belo": "Belo", "fwind": "Fiwind", "fiwind": "Fiwind", "astropay": "AstroPay", "banco galicia": "Banco Galicia",
    "lemon cash": "Lemon Cash", "lemon": "Lemon Cash", "mercado pago": "Mercado Pago", "naraja x": "Naranja X",
    "naranja x": "Naranja X", "ripio": "Ripio", "dolar app": "DolarApp",
}


def _num(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def _sin_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _norm(x, tabla):
    if not isinstance(x, str) or not x.strip():
        return ""
    k = " ".join(x.strip().lower().split())
    if "mariano" in k:
        return "Cliente del referente (Mariano)"
    return tabla.get(k, x.strip() if not x.strip().islower() else x.strip().title())


def es_mensual(ws):
    return ws.cell(1, 2).value == "ID VENTA" and ws.cell(1, 12).value == "ID COMPRA"


def mes_de_nombre(nombre):
    """Número de mes a partir del nombre de la hoja ('Registro USDT General-(SEP)' → 9)."""
    s = _sin_tildes(nombre.lower()).replace("registro usdt general", " ").replace("coco", " ")
    for pal in re.findall(r"[a-z]+", s):
        if pal[:3] in MESES:
            return MESES[pal[:3]]
    raise ValueError(f"No se reconoce el mes de la hoja {nombre!r}")


def _limites(anio, mes):
    ini = dt.date(anio, mes, 1)
    return ini, dt.date(anio, mes, calendar.monthrange(anio, mes)[1])


def _fecha(x):
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    return None


def _anio_hoja(ws, mes):
    """Año más frecuente de las fechas del mes de la hoja (columnas A y K)."""
    c = Counter()
    for col in (1, 11):
        for r in range(3, ws.max_row + 1):
            d = _fecha(ws.cell(r, col).value)
            if d and d.month == mes:
                c[d.year] += 1
    return c.most_common(1)[0][0] if c else None


def _corregir(d, prev, sig, ini, fin):
    """Fecha de la celda, corregida si cae lejos del mes de la hoja. Devuelve (fecha, corregida)."""
    if ini - dt.timedelta(days=VENTANA_DIAS) <= d <= fin + dt.timedelta(days=VENTANA_DIAS):
        return d, False
    try:
        cand = dt.date(ini.year, ini.month, d.day)
    except ValueError:
        cand = None
    if cand and (prev is None or prev <= cand) and (sig is None or cand <= sig):
        return cand, True
    return (prev, True) if prev else (None, True)


class Libro:
    """Todo lo que el tablero necesita del Excel, ya normalizado."""

    def __init__(self, ruta):
        self.ruta = ruta
        wv = openpyxl.load_workbook(ruta, data_only=True)
        wf = openpyxl.load_workbook(ruta, read_only=True)
        self.hojas = []      # [{nombre, libro, anio, mes, clave}]
        self.filas = []      # filas de datos (dict)
        self.notas_ws = {}
        for ws in wv.worksheets:
            if es_mensual(ws):
                self._leer_mensual(ws, wf[ws.title])
        self.hojas.sort(key=lambda h: (h["clave"], h["libro"] != "principal"))
        self.resultado = self._leer_resultado(wv["RESULTADO USDT GENERAL"]) if "RESULTADO USDT GENERAL" in wv.sheetnames else []
        self.referente = self._leer_referente(wv["CLIENTES MARIANO"]) if "CLIENTES MARIANO" in wv.sheetnames else []
        fechas = [f["fecha"] for f in self.filas if f["fecha"]]
        self.corte = max(fechas) if fechas else None

    # ------------------------------------------------------------------ hojas mensuales
    def _leer_mensual(self, ws, wsf):
        mes = mes_de_nombre(ws.title)
        anio = _anio_hoja(ws, mes)
        libro = "coco" if "COCO" in ws.title.upper() else "principal"
        clave = f"{anio}-{mes:02d}"
        self.hojas.append({"nombre": ws.title, "libro": libro, "anio": anio, "mes": mes, "clave": clave})
        ini, fin = _limites(anio, mes)
        formulas = {}
        for fila in wsf.iter_rows(min_row=3, max_col=19):
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas[(c.row, c.column)] = c.value
        ult = ws.max_row
        for lado, (cf, ci, cb, cc, ca, cx, cbk) in LADOS.items():
            celdas = {r: _fecha(ws.cell(r, cf).value) for r in range(3, ult + 1)}
            con_fecha = [r for r in sorted(celdas) if celdas[r]]
            fecha, prev = None, None
            corr_por_fila = {}
            for k, r in enumerate(con_fecha):
                sig = None
                for r2 in con_fecha[k + 1:]:
                    d2 = celdas[r2]
                    if ini - dt.timedelta(days=VENTANA_DIAS) <= d2 <= fin + dt.timedelta(days=VENTANA_DIAS):
                        sig = d2
                        break
                d, corr = _corregir(celdas[r], prev, sig, ini, fin)
                corr_por_fila[r] = (d, corr, celdas[r])
                if d:
                    prev = d
            fecha_corr = False
            for r in range(3, ult + 1):
                if r in corr_por_fila:
                    fecha, fecha_corr, cargada = corr_por_fila[r]
                bruto = _num(ws.cell(r, cb).value)
                ars = _num(ws.cell(r, ca).value)
                idv = ws.cell(r, ci).value
                usdt_del_id = False
                if not bruto and ars and _num(idv) and _num(idv) > 0 and PRECIO_ID[0] <= ars / _num(idv) <= PRECIO_ID[1]:
                    bruto, usdt_del_id = _num(idv), True
                if not bruto and not ars:
                    continue
                com = _num(ws.cell(r, cc).value) or 0.0
                bruto = bruto or 0.0
                usdt = bruto + com if lado == "V" else bruto - com
                if not bruto:
                    usdt = 0.0
                id_txt = "" if usdt_del_id or idv is None else (str(idv) if not isinstance(idv, float) else f"{idv:.0f}")
                avisos = []
                if usdt_del_id:
                    avisos.append("usdt_del_id")
                if not usdt or not ars:
                    avisos.append("incompleta")
                col_neto = 5 if lado == "V" else 15
                if bruto and (r, col_neto) not in formulas and _num(ws.cell(r, col_neto).value) is None:
                    avisos.append("sin_formula_neto")
                if fecha_corr:
                    avisos.append("fecha_corregida")
                if fecha is None:
                    avisos.append("sin_fecha")
                self.filas.append({
                    "hoja": ws.title, "libro": libro, "mes": clave, "lado": lado, "fila": r,
                    "fecha": fecha, "id": id_txt, "bruto": bruto, "comision": com, "usdt": usdt, "ars": ars or 0.0,
                    "canal_raw": ws.cell(r, cx).value, "canal": _norm(ws.cell(r, cx).value, CANALES) or "Sin canal",
                    "banco": _norm(ws.cell(r, cbk).value, BANCOS) or "Sin cuenta",
                    "precio": (ars / usdt) if usdt and ars else None, "avisos": avisos,
                })
        # Lo que dice la propia hoja (fila 2), para comparar
        g = lambda c: _num(ws.cell(2, c).value)
        self.hojas[-1]["hoja_dice"] = {"E2": g(5), "O2": g(15), "G2": g(7), "Q2": g(17), "U2": g(21), "V2": g(22)}
        self._marcar_precios(ws.title)

    def _marcar_precios(self, hoja):
        for lado in "VC":
            fs = [f for f in self.filas if f["hoja"] == hoja and f["lado"] == lado]
            for i, f in enumerate(fs):
                if f["precio"] is None:
                    continue
                vec = [x["precio"] for x in fs[max(0, i - 8):i + 8] if x is not f and x["precio"]]
                if len(vec) >= 5:
                    med = statistics.median(vec)
                    if abs(f["precio"] / med - 1) > 0.04:
                        f["avisos"].append("precio_raro")
                        f["precio_vecinos"] = med

    # ------------------------------------------------------------------ RESULTADO
    def _leer_resultado(self, ws):
        bloques, anio_prev = [], None
        filas = {r: [ws.cell(r, c).value for c in range(1, 13)] for r in range(1, ws.max_row + 1)}
        for r, v in filas.items():
            if not (isinstance(v[0], str) and v[0].strip().upper() == "MES"):
                continue
            titulo, libro = None, "principal"
            for r2 in range(r - 1, max(0, r - 4), -1):
                t = filas.get(r2, [None])[0]
                if isinstance(t, str) and t.strip().upper().startswith("BALANCE"):
                    titulo = t
                if isinstance(t, str) and t.strip().upper() == "COCO":
                    libro = "coco"
            anio = int(re.findall(r"\d{4}", titulo)[0]) if titulo else (anio_prev + 1 if anio_prev else None)
            if libro == "principal":
                anio_prev = anio
            nombres = [x.strip() if isinstance(x, str) else "" for x in filas[r + 1][5:11]]
            meses = []
            for r2 in range(r + 2, r + 16):
                a = filas.get(r2, [None])[0]
                if not isinstance(a, str):
                    continue
                if a.strip().upper() == "TOTAL":
                    break
                m = MESES.get(_sin_tildes(a.strip().lower())[:3])
                if not m:
                    continue
                v2 = [(_num(x) or 0.0) for x in filas[r2][1:11]]
                meses.append({"clave": f"{anio}-{m:02d}", "fila": r2, "bruto": v2[0], "impuestos": v2[1],
                              "gastos": v2[2], "neto": v2[3], "partes": v2[4:7], "retiros": v2[7:10]})
            bloques.append({"anio": anio, "libro": libro, "fila": r, "socios": nombres[:3],
                            "retiran": nombres[3:6], "meses": meses})
        return bloques

    # ------------------------------------------------------------------ CLIENTES MARIANO
    def _leer_referente(self, ws):
        out, fecha = [], None
        fechas = {r: _fecha(ws.cell(r, 2).value) for r in range(3, ws.max_row + 1)}
        con = [r for r in sorted(fechas) if fechas[r]]
        for k, r in enumerate(con):
            d = fechas[r]
            prev = fechas[con[k - 1]] if k else None
            sig = fechas[con[k + 1]] if k + 1 < len(con) else None
            if prev and sig and prev <= sig and not (prev <= d <= sig):
                try:
                    cand = dt.date(prev.year, prev.month, d.day)
                    if prev <= cand <= sig:
                        fechas[r] = cand
                except ValueError:
                    pass
        for r in range(3, ws.max_row + 1):
            if fechas.get(r):
                fecha = fechas[r]
            j = _num(ws.cell(r, 10).value)
            if j is None or fecha is None:
                continue
            out.append({"fila": r, "fecha": fecha, "mes": f"{fecha.year}-{fecha.month:02d}", "ganancia": j,
                        "ars": _num(ws.cell(r, 4).value), "usdt": _num(ws.cell(r, 7).value)})
        return out
