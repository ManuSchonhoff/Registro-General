"""Salidas: Excel de análisis, metricas.json y resumen en Markdown para el asesor."""
import json
import math

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3864")
B = Font(name=F)
T = Font(name=F, bold=True, size=14)


def _fmt(col):
    c = str(col).lower()
    if "pactado" in c or c.endswith("_pct"):
        return '0.00"%"'  # ya viene expresado en % (2,5 = 2,5%)
    if "%" in c or "pct" in c or "margen" in c or "cobertura" in c or "tasa" in c:
        return "0.0%"
    if "cot" in c or "spread_pesos" in c:
        return "#,##0.00"
    if any(k in c for k in ("usd", "monto", "ars", "debe", "haber", "volumen", "ganancia", "comision", "saldo", "gasto",
                            "interes", "total", "neto", "efecto")):
        return "#,##0;-#,##0;-"
    return None


def _limpio(v):
    if isinstance(v, (np.floating, float)):
        return None if (math.isnan(v) or math.isinf(v)) else float(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (pd.Timestamp,)):
        return v.to_pydatetime()
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def excel(tablas, notas, ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Índice"
    ws["A1"] = "Análisis del Registro General"
    ws["A1"].font = T
    r = 3
    for k, v in notas.items():
        ws.cell(r, 1, k).font = Font(name=F, bold=True)
        ws.cell(r, 2, v).font = B
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(r, 1, "Hojas").font = Font(name=F, bold=True)
    for nombre in tablas:
        r += 1
        ws.cell(r, 1, nombre[:31]).font = B
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    for nombre, df in tablas.items():
        if df is None or df.empty:
            continue
        sh = wb.create_sheet(nombre[:31])
        df = df.reset_index(drop=True)
        for j, col in enumerate(df.columns, 1):
            c = sh.cell(1, j, str(col))
            c.font, c.fill = H, HF
            c.alignment = Alignment(wrap_text=True, vertical="center")
            fmt = _fmt(col)
            ancho = max(10, min(40, len(str(col)) + 2))
            for i, v in enumerate(df[col].tolist(), 2):
                cell = sh.cell(i, j, _limpio(v))
                cell.font = B
                if fmt and isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
                if isinstance(v, str):
                    ancho = max(ancho, min(45, len(v) + 2))
            sh.column_dimensions[get_column_letter(j)].width = ancho
        sh.freeze_panes = "B2"
        sh.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    wb.save(ruta)


def _num(v):
    if isinstance(v, float):
        if abs(v) >= 1000:
            return f"{v:,.0f}".replace(",", ".")
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


def md_tabla(df, max_filas=15):
    df = df.head(max_filas)
    cols = [str(c) for c in df.columns]
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for row in df.itertuples(index=False):
        celdas = []
        for c, v in zip(cols, row):
            if isinstance(v, (float, np.floating)) and not math.isnan(v):
                ya_en_pct = "pactado" in c or c.endswith("_pct")
                es_frac = ("%" in c or "pct" in c or "margen" in c or "tasa" in c) and not ya_en_pct and abs(v) < 5
                celdas.append(f"{v:.1%}" if es_frac else (f"{v:.2f}%" if ya_en_pct else _num(float(v))))
            else:
                celdas.append("" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v))
        out.append("| " + " | ".join(celdas) + " |")
    return "\n".join(out)


def metricas_json(metricas, ruta):
    def conv(o):
        if isinstance(o, dict):
            return {str(k): conv(v) for k, v in o.items()}
        if isinstance(o, list):
            return [conv(v) for v in o]
        return _limpio(o)
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(conv(metricas), f, ensure_ascii=False, indent=1, default=str)


def resumen_md(tablas, metricas, contexto, ruta):
    """Documento compacto que lee el agente asesor: métricas clave + tablas principales."""
    secciones = [
        ("Volumen y clientes", ["01 Ranking clientes", "02 Altas y bajas (mes)", "02 Clientes perdidos", "03 Operador resumen"]),
        ("Rentabilidad", ["Resultado estimado mensual", "04 Spread mensual", "04 Spread por operador", "05 Circuito mensual",
                          "05 Circuito por operador", "06 Comisiones por tipo", "07 Cheques por titular"]),
        ("Operación", ["08 Evolución mensual", "09 Día de la semana", "10 Tamaño ops", "10 Ticket por operador",
                       "11 Calidad por operador"]),
        ("Posición, pendientes y fondeo", ["12 Posición USD (saldo real)", "12 Saldo vs caja real", "12 Inversiones y garantías", "13 Pendientes antigüedad",
                                           "13 Pendientes por cliente", "14 Inversores", "14 Fondeo vs ingresos",
                                           "15 Gastos por categoría"]),
    ]
    L = ["# Resumen de datos para el asesor financiero", ""]
    for k, v in contexto.items():
        L.append(f"- **{k}:** {v}")
    L += ["", "## Métricas clave (metricas.json)", "", "```json",
          json.dumps({k: v for k, v in metricas.items()}, ensure_ascii=False, indent=1, default=str)[:60000], "```", ""]
    for titulo, hojas in secciones:
        L += [f"## {titulo}", ""]
        for h in hojas:
            df = tablas.get(h)
            if df is None or df.empty:
                continue
            L += [f"### {h}", "", md_tabla(df.reset_index(drop=True), 25 if "Resultado" in h or "mensual" in h else 15), ""]
    with open(ruta, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
