"""Flujo semanal de un paso del tablero (docs/tablero_v2_spec.md, sección 7).

Lo usa actualizar.py (la línea de comando, en la raíz):
- planificar: qué Excel de datos/ es la base, cuáles son actualizaciones y cuáles se ignoran, y qué archivo aporta
  cada semana. Sale de los nombres de archivo y de hoja (lectura read_only, rápida).
- ultima_semana / fecha_corte: lee solo la última hoja para saber si está cerrada y la fecha del último Cierre.
- chequeos_plan: los chequeos que se pueden evaluar sin procesar (C01 liviano, C05, C06 y C16); C02–C04 quedan
  "no_evaluado" hasta que exista el escaneo completo (P1).
- notas: salidas/notas_manuales.json (cargar, migrar, sembrar) y la copia de las notas abiertas para el v1
  (arrastrar_v1).
- publicación: rutas fijas salidas/tablero_v1_artifact.html y salidas/tablero_v2_artifact.html, y
  salidas/estado_publicacion.json (Python nunca borra una url).

P1 (todavía no): escanear el contenido de cada hoja (cerrada, Ope máximo, Cierres, fórmulas sin valor) y pasar
`elegir`/`excluir` a Registro. Mientras tanto, la base va primero y las actualizaciones en orden de fecha.
"""
import datetime as dt
import hashlib
import json
import os
import re
import shutil

import openpyxl

from registro import tablero_v2 as TB2
from registro.leer import SEMANA, Semana

NOTAS = "notas_manuales.json"
NOTAS_V1 = "notas_tablero.json"
NOTAS_V1_ORIGINAL = "notas_tablero_original.json"
DATOS_V2 = "tablero_v2_datos.json"
ESTADO = "estado_publicacion.json"
RUTAS_FIJAS = {"v1": "tablero_v1_artifact.html", "v2": "tablero_v2_artifact.html"}
CAMPOS_ESTADO = ("url", "fecha_corte", "semana", "semanas", "sha256", "generado")
NOTA_SEMBRADO = "error de carga ya resuelto"
MARCA_V1 = "_nota"
TEXTO_MARCA_V1 = ("Copia generada por actualizar.py con las notas abiertas de salidas/notas_manuales.json. "
                  "Las notas se editan allá, no acá.")

COPIA = re.compile(r"^(copia de\s+)+", re.I)
HOJA_SUELTA = re.compile(r"^\d{4}\s*-\s*\d{4}")
FECHA_NOMBRE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
CARPETA_FECHA = re.compile(r"\d{4}-\d{2}-\d{2}")


class ErrorUsuario(Exception):
    """Error de uso o de archivos (se muestra el mensaje solo, sin el detalle técnico)."""


# ============================================================================ auxiliares
def _leer_json(ruta):
    with open(ruta, encoding="utf-8") as fh:
        return json.load(fh)


def _escribir_json(ruta, obj, indent=2):
    """Escritura atómica: si se corta a la mitad, el archivo anterior queda intacto."""
    os.makedirs(os.path.dirname(os.path.abspath(ruta)), exist_ok=True)
    tmp = ruta + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=indent)
        fh.write("\n")
    os.replace(tmp, ruta)


def sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as fh:
        for bloque in iter(lambda: fh.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def nombre_semana(titulo):
    """Nombre con el que Registro lee la hoja (sin "Copia de"), o None si no es semanal. Mismo criterio que leer.py."""
    n = COPIA.sub("", str(titulo).strip())
    return n if SEMANA.match(n) else None


def fechas_semana(nombre, anio):
    """(inicio, fin) de una hoja "DDMM - DDMM", con la misma regla que leer.Semana (el fin puede caer en el año siguiente)."""
    m = SEMANA.match(nombre)
    ini = dt.date(anio, int(m.group(2)), int(m.group(1)))
    fin = dt.date(anio, int(m.group(4)), int(m.group(3)))
    if fin < ini:
        fin = dt.date(anio + 1, fin.month, fin.day)
    return ini, fin


def _orden_semana(nombre):
    m = SEMANA.match(nombre)
    return int(m.group(2)), int(m.group(1))  # el mismo orden que Registro: (mes, día) de inicio


def etiqueta(nombre, anio):
    ini, fin = fechas_semana(nombre, anio)
    return TB2.etiqueta_semana(ini, fin)


# ============================================================================ plan (7.1, paso 1)
def _hojas(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def listar_excels(carpeta):
    if not os.path.isdir(carpeta):
        raise ErrorUsuario(f"No existe la carpeta {carpeta}")
    return sorted(os.path.join(carpeta, f) for f in os.listdir(carpeta)
                  if f.lower().endswith(".xlsx") and not f.startswith(("~$", ".")))


def planificar(archivos, cfg):
    """Plan P0 (7.1). Base: el archivo más reciente con las hojas CQ e Inversores (va primero, porque Registro toma de
    ahí las hojas auxiliares). Actualizaciones: los posteriores a la base, en orden de fecha. Los anteriores a la base se
    ignoran (aplicados después pisarían hojas más nuevas).

    Fecha de cada archivo: la AAAA-MM-DD del nombre; si no tiene, el fin de su última hoja semanal; la fecha de
    modificación solo desempata (cambia cada vez que se vuelve a subir el archivo)."""
    anio = int(cfg["anio"])
    if not archivos:
        raise ErrorUsuario("No hay archivos .xlsx para procesar")
    info = []
    for ruta in archivos:
        hojas = _hojas(ruta)
        semanas, sueltas, otras = [], [], []
        for h in hojas:
            n = nombre_semana(h)
            if n:
                try:
                    fechas_semana(n, anio)
                except ValueError:
                    sueltas.append(h)
                    continue
                semanas.append((h, n))
            elif HOJA_SUELTA.match(COPIA.sub("", h.strip())):
                sueltas.append(h)
            else:
                otras.append(h)
        nombre = os.path.basename(ruta)
        m = FECHA_NOMBRE.search(nombre)
        fecha, origen = None, None
        if m:
            try:
                fecha, origen = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))), "nombre"
            except ValueError:
                fecha = None
        if fecha is None and semanas:
            fecha, origen = max(fechas_semana(n, anio)[1] for _, n in semanas), "hojas"
        mtime = os.path.getmtime(ruta)
        if fecha is None:
            fecha, origen = dt.date.fromtimestamp(mtime), "modificación"
        info.append(dict(ruta=ruta, archivo=nombre, fecha=fecha.isoformat(), fecha_origen=origen, mtime=mtime,
                         cq="CQ" in hojas, inversores="Inversores" in hojas, semanas=semanas, sueltas=sueltas,
                         otras=otras, rol=None, motivo=None))
    if len({d["archivo"] for d in info}) < len(info):
        raise ErrorUsuario("Hay dos archivos con el mismo nombre en carpetas distintas: pasá uno solo de cada uno")

    def orden(d):
        return d["fecha"], d["mtime"], d["archivo"]

    con_aux = [d for d in info if d["cq"] and d["inversores"]]
    base = max(con_aux or info, key=orden)
    base["rol"] = "base"
    if not con_aux:
        base["motivo"] = "ningún archivo tiene las hojas CQ e Inversores"
    for d in info:
        if d is base:
            continue
        if orden(d) > orden(base):
            d["rol"] = "actualizacion"
            if d["fecha"] == base["fecha"]:
                d["motivo"] = "misma fecha que la base: se ordenó por fecha de modificación"
        else:
            d["rol"] = "ignorado"
            d["motivo"] = "es anterior a la base (aplicado después pisaría hojas más nuevas)"
        if d["rol"] == "actualizacion" and not d["semanas"]:
            d["rol"], d["motivo"] = "ignorado", "no tiene hojas semanales"
    usados = [base] + sorted((d for d in info if d["rol"] == "actualizacion"), key=orden)
    semana_archivo, semana_hoja = {}, {}
    for d in usados:
        for h, n in d["semanas"]:
            semana_archivo[n], semana_hoja[n] = d["archivo"], h  # como Registro: la última aplicada gana
    nombres = sorted(semana_archivo, key=_orden_semana)
    if not nombres:
        raise ErrorUsuario("Ningún archivo tiene hojas semanales (\"DDMM - DDMM\")")
    return dict(archivos=sorted(info, key=orden), rutas=[d["ruta"] for d in usados], base=base["ruta"],
                base_ok=bool(con_aux), semanas=nombres, semana_archivo=semana_archivo, semana_hoja=semana_hoja,
                anio=anio)


def escanear(plan):
    """Escaneo P0 (liviano): semana → archivo según el orden del plan, sin leer el contenido. El escaneo P1 agrega, por
    hoja y archivo, si está cerrada, el Ope máximo, si tiene Cierres y el % de fórmulas sin valor."""
    return [dict(semana=n, archivo=plan["semana_archivo"][n], hoja=plan["semana_hoja"][n]) for n in plan["semanas"]]


# ============================================================================ fecha de corte (7.1, paso 3)
def ultima_semana(plan):
    """Lee solo la hoja de la última semana (del archivo que la aporta) para saber si está cerrada y la fecha del
    último Cierre con fecha válida en la columna Q."""
    nombre = plan["semanas"][-1]
    archivo = plan["semana_archivo"][nombre]
    ruta = next(d["ruta"] for d in plan["archivos"] if d["archivo"] == archivo)
    hoja = plan["semana_hoja"][nombre]
    wf = openpyxl.load_workbook(ruta)
    wv = openpyxl.load_workbook(ruta, data_only=True)
    s = Semana(nombre, wf[hoja], wv[hoja], plan["anio"], len(plan["semanas"]))
    fechas = []
    for r in s.filas("Cierre"):
        q = s.v.cell(r, 17).value
        if isinstance(q, dt.datetime) and s.inicio <= q.date() <= s.fin:
            fechas.append(q.date())
    return dict(nombre=nombre, ini=s.inicio.isoformat(), fin=s.fin.isoformat(), archivo=archivo, cerrada=s.cerrada(),
                ultimo_cierre=max(fechas).isoformat() if fechas else None)


def fecha_corte(ultima, hoy=None):
    """Si la última semana está cerrada, su domingo; si no, el último Cierre con fecha o, si no hay, hoy, sin pasar el
    fin de la semana. Devuelve (fecha, motivo)."""
    if ultima["cerrada"]:
        return ultima["fin"], "la última semana está cerrada: su domingo"
    if ultima["ultimo_cierre"]:
        return ultima["ultimo_cierre"], "semana en curso: el último Cierre con fecha"
    hoy = (hoy or dt.date.today()).isoformat()
    return min(max(hoy, ultima["ini"]), ultima["fin"]), "semana en curso sin Cierres con fecha: hoy"


def ultima_cerrada_ini(plan, ultima):
    """ini de la última semana cerrada (la foco por defecto), o None si no hay ninguna."""
    if ultima["cerrada"]:
        return ultima["ini"]
    if len(plan["semanas"]) < 2:
        return None
    return fechas_semana(plan["semanas"][-2], plan["anio"])[0].isoformat()


# ============================================================================ chequeos del plan (7.2)
def _check(id_, nivel, detalle, claves=(), n=None):
    claves = list(claves)
    return dict(id=id_, nivel=nivel, titulo=TB2.TITULOS[id_], detalle=detalle, claves=claves,
                n=len(claves) if n is None else n)


def chequeos_escaneo(plan):
    """C01–C04, para plan["control"] (el pipeline no los evalúa). C01 sale de los nombres de hoja; C02–C04 necesitan el
    escaneo completo (P1) y quedan como no evaluados."""
    sueltas = [(d["archivo"], h) for d in plan["archivos"] if d["rol"] != "ignorado" for h in d["sueltas"]]
    extra = [(d["archivo"], h) for d in plan["archivos"] if d["rol"] == "actualizacion" for h in d["otras"]]
    out = []
    if sueltas or extra:
        partes = []
        if sueltas:
            partes.append("parecen semanales pero el nombre no tiene la forma de siempre (día y mes de inicio y de fin, "
                          "como «1210 - 1810»), así que no se leen: "
                          + ", ".join(f"«{h}» ({a})" for a, h in sueltas))
        if extra:
            partes.append("no son semanales y vienen en un archivo de actualización: "
                          + ", ".join(f"«{h}» ({a})" for a, h in extra))
        out.append(_check("C01", "aviso", "Hojas que " + "; ".join(partes) + ".",
                          [f"control:C01:{h}" for _, h in sueltas + extra]))
    else:
        out.append(_check("C01", "ok", "Todas las hojas semanales tienen nombre reconocible.", n=0))
    pend = "Chequeo todavía no disponible."
    for cid in ("C02", "C03", "C04"):
        out.append(_check(cid, "no_evaluado", pend, n=0))
    return out


def chequeos_plan(plan, previo, fecha, cfg):
    """Lo que --dry-run puede decir antes de procesar: C01–C04 (como en chequeos_escaneo), C05, C06 y C16. En la
    corrida completa C05, C06 y C16 los vuelve a evaluar el pipeline con la planilla leída."""
    out = chequeos_escaneo(plan)
    base = next(d for d in plan["archivos"] if d["rol"] == "base")
    faltan = [h for h, ok in (("CQ", base["cq"]), ("Inversores", base["inversores"])) if not ok]
    if faltan:
        out.append(_check("C05", "bloqueante", f"La base {base['archivo']} no tiene la hoja {' ni la hoja '.join(faltan)}.",
                          [f"control:C05:{base['archivo']}"]))
    else:
        out.append(_check("C05", "ok", f"La base {base['archivo']} tiene las hojas CQ e Inversores.", n=0))
    n = len(plan["semanas"])
    prev_n = None
    if isinstance(previo, dict):
        prev_n = (previo.get("v2") or {}).get("semanas") or (previo.get("v1") or {}).get("semanas")
    if prev_n is None:
        out.append(_check("C06", "no_evaluado", "No hay una publicación anterior para comparar.", n=0))
    elif n < int(prev_n):
        out.append(_check("C06", "bloqueante", f"Hay {n} {'semana' if n == 1 else 'semanas'} y la última publicación "
                          f"tenía {prev_n}: falta alguna hoja.", [f"control:C06:{n}"]))
    else:
        out.append(_check("C06", "ok", f"{n} semanas (la última publicación tenía {prev_n}).", n=0))
    anio = int(cfg["anio"])
    cruza = [s for s in plan["semanas"] if fechas_semana(s, anio)[1].year != anio]
    corte = dt.date.fromisoformat(fecha).year
    if anio != corte or cruza:
        motivo = (f"la configuración del tablero dice {anio} y la fecha de corte es de {corte}" if anio != corte
                  else f"la hoja {cruza[0]} cruza el año")
        out.append(_check("C16", "aviso", f"Revisar el año: {motivo}.", [f"control:C16:{anio}"]))
    else:
        out.append(_check("C16", "ok", f"Todas las hojas son de {anio}.", n=0))
    return out


def validar(R, b, T, M, cfg, plan, previo):
    """Chequeos de 7.2 desde C05 con la planilla ya leída (delegado en tablero_v2.validar). actualizar.py no lo
    necesita porque el pipeline ya los calcula; queda para usarlo desde otro script."""
    return TB2.validar(R, b, T, M, cfg, plan=plan, previo=previo)


# ============================================================================ notas (7.1 paso 5, 7.3)
def _carpetas_fecha(raiz):
    if not os.path.isdir(raiz):
        return []
    return sorted(d for d in os.listdir(raiz) if CARPETA_FECHA.fullmatch(d) and os.path.isdir(os.path.join(raiz, d)))


def _cuantas(notas):
    return len(notas["hallazgos"]) + len(notas["preguntas"])


def leer_paquete(ruta):
    """Datos de un tablero v2: tablero_v2_datos.json, o el HTML (o el Artifact guardado) sin clave."""
    with open(ruta, encoding="utf-8") as fh:
        txt = fh.read()
    if ruta.lower().endswith(".json"):
        d = json.loads(txt)
    else:
        m = re.search(r"const\s+PAQUETE\s*=\s*", txt)
        if not m:
            raise ErrorUsuario(f"{ruta} no parece un tablero v2 (no tiene «const PAQUETE =»)")
        d, _ = json.JSONDecoder().raw_decode(txt, m.end())
        if d is None:
            raise ErrorUsuario(f"{ruta} es la plantilla sin datos")
    if isinstance(d, dict) and "plano" in d:
        d = d["plano"]
    elif isinstance(d, dict) and "ct" in d:
        raise ErrorUsuario(f"{ruta} tiene clave: no se pueden leer sus notas")
    return d


def cargar(raiz, semanas, desde=None):
    """salidas/notas_manuales.json si existe. Si no, se reconstruye (sin escribir) desde: `desde` (un tablero v2
    publicado, guardado localmente), las notas del último salidas/*/tablero_v2_datos.json o, si no hay, migrando el
    notas_tablero.json más reciente. semanas: lista de (inicio, fin) para ubicar la semana de cada hallazgo.
    Devuelve (notas, origen) con origen = {"tipo": "manuales" | "migradas" | "ninguna", "archivo", "n"}."""
    man = os.path.join(raiz, NOTAS)
    if os.path.exists(man):
        return TB2.normalizar_notas(_leer_json(man)), dict(tipo="manuales", archivo=man, n=0)
    if desde:
        d = leer_paquete(desde)
        notas = TB2.normalizar_notas(d.get("notas") if "notas" in d else d)
        return notas, dict(tipo="migradas", archivo=os.path.abspath(desde), n=_cuantas(notas))
    for c in reversed(_carpetas_fecha(raiz)):
        p = os.path.join(raiz, c, DATOS_V2)
        if os.path.exists(p):
            d = _leer_json(p)
            if isinstance(d.get("notas"), dict):
                notas = TB2.normalizar_notas(d["notas"])
                return notas, dict(tipo="migradas", archivo=p, n=_cuantas(notas))
    for c in reversed(_carpetas_fecha(raiz)):
        p = os.path.join(raiz, c, NOTAS_V1)
        if os.path.exists(p):
            notas, n = migrar(_leer_json(p), c, semanas)
            return notas, dict(tipo="migradas", archivo=p, n=n)
    return TB2.normalizar_notas({}), dict(tipo="ninguna", archivo=None, n=0)


def migrar(v1, fecha_origen, semanas, existentes=None):
    """notas_tablero.json del v1 → formato de notas_manuales.json, sin duplicar textos (7.3)."""
    return TB2.migrar_notas(v1, fecha_origen, semanas, existentes)


def absorber_v1(notas, salida, fecha, semanas, origen=None, raiz=None):
    """Si salidas/<fecha>/notas_tablero.json fue escrito a mano (no lo generó actualizar.py) DESPUÉS del último guardado
    de notas_manuales.json (o si notas_manuales.json todavía no existe), suma a `notas` los textos que todavía no están,
    para no perderlos al reemplazarlo por la copia de las notas abiertas. Modifica `notas`.

    Un notas_tablero.json escrito a mano ANTES del último guardado (por ejemplo, el de una carpeta de antes del v2) no se
    absorbe: sus notas ya pasaron por notas_manuales.json y pueden estar resueltas o reemplazadas; sumarlas las
    reabriría. Ese archivo lo guarda arrastrar_v1 una vez como notas_tablero_original.json y lo reemplaza por la copia.

    Devuelve (textos agregados, motivo): motivo es None si no había nada para absorber, "absorbido" o "viejo"
    (escrito a mano antes del último guardado: no se suma nada)."""
    p = os.path.join(salida, NOTAS_V1)
    if not os.path.exists(p):
        return [], None
    v1 = _leer_json(p)
    if not isinstance(v1, dict) or v1.get(MARCA_V1):
        return [], None
    man = os.path.join(raiz, NOTAS) if raiz else None
    hay_manuales = (origen or {}).get("tipo") == "manuales" and man and os.path.exists(man)
    if hay_manuales and os.path.getmtime(p) <= os.path.getmtime(man):
        return [], "viejo"
    antes = {(k, x["texto"]) for k in ("hallazgos", "preguntas") for x in notas[k]}
    nuevas, n = migrar(v1, fecha, semanas, existentes=notas)
    if not n:
        return [], None
    notas.clear()
    notas.update(nuevas)
    textos = [x["texto"] for k in ("hallazgos", "preguntas") for x in notas[k] if (k, x["texto"]) not in antes]
    return textos, "absorbido"


def guardar_notas(raiz, notas):
    """Escribe notas_manuales.json con una entrada por línea, para buscar y editar a mano una clave o una nota."""
    n = TB2.normalizar_notas(notas)
    partes = []
    for k in ("hallazgos", "preguntas", "revisadas"):
        filas = ",\n".join("  " + json.dumps(x, ensure_ascii=False) for x in n[k])
        partes.append(f'"{k}": [\n{filas}\n ]' if filas else f'"{k}": []')
    ruta = os.path.join(raiz, NOTAS)
    os.makedirs(raiz, exist_ok=True)
    tmp = ruta + ".tmp"
    texto = "{" + ",\n ".join(partes) + "}\n"
    if TB2.normalizar_notas(json.loads(texto)) != n:  # nunca dejar un JSON roto o distinto de lo que se quiso guardar
        raise ValueError("Error al armar notas_manuales.json")
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write(texto)
    os.replace(tmp, ruta)


def ya_sembrado(notas):
    """Fecha del sembrado anterior (alguna revisada con la nota del sembrado), o None."""
    return next((e.get("fecha") or "?" for e in notas["revisadas"] if e.get("nota") == NOTA_SEMBRADO), None)


def tope_sembrado(cfg, hasta_ini):
    """ini de la última semana que --sembrar-revisadas puede dar por revisada: la última cerrada, pero nunca después de
    tablero_v2.sembrar_revisadas_hasta (lo que el dueño confirmó el 19/10/2026). Así, si se pierden las notas y se vuelve
    a sembrar más adelante, no se dan por revisadas filas que nadie confirmó."""
    tope = TB2.cfg_v2(cfg).get("sembrar_revisadas_hasta")
    if not tope or not hasta_ini:
        return hasta_ini
    # la semana que contiene el tope: su lunes
    t = dt.date.fromisoformat(tope)
    lunes = (t - dt.timedelta(days=t.weekday())).isoformat()
    return min(hasta_ini, lunes)


def sembrar_hallazgos(notas, hasta_ini, hoy):
    """--sembrar-revisadas: los hallazgos abiertos con semana ≤ la última cerrada (o sin semana) pasan a resueltos."""
    n = 0
    for h in notas["hallazgos"]:
        if h.get("estado") == "abierta" and (h.get("semana") is None or (hasta_ini and h["semana"] <= hasta_ini)):
            h["estado"], h["fecha_estado"] = "resuelta", hoy
            n += 1
    return n


def claves_revisar(R, b, T, hasta_ini=None):
    """Claves <ini>#<fila>#<ope> de las filas Alta y Media de las semanas cerradas (las mismas de `revisar`), con
    inicio ≤ hasta_ini si se indica (tope del sembrado)."""
    t = T.get("11 Filas a revisar")
    if t is None or t.empty:
        return []
    t = t[t.severidad.isin(["Alta", "Media"])]
    base = b.ops[["semana_n", "fila", "ope"]].drop_duplicates(["semana_n", "fila"])
    t = t.merge(base, on=["semana_n", "fila"], how="left", suffixes=("_t", ""))
    ultima_cerrada = len(R.semanas) - (1 if R.actual_cerrada else 2)
    out = []
    for r in t.sort_values(["semana_n", "fila"]).itertuples():
        w = int(r.semana_n) - 1
        if 0 <= w <= ultima_cerrada and (hasta_ini is None or R.semanas[w].inicio.isoformat() <= hasta_ini):
            out.append(TB2.clave_fila(R.semanas[w].inicio.isoformat(), r.fila, r.ope))
    return out


def sembrar_filas(notas, claves, hoy):
    """Agrega a `revisadas` las claves que todavía no están (por coincidencia de 7.3). Devuelve cuántas agregó."""
    rev = TB2.Revisadas(notas["revisadas"])
    vistas, n = set(), 0
    for k in claves:
        if k in vistas or rev.clave(k) is not None:
            continue
        vistas.add(k)
        notas["revisadas"].append(dict(clave=k, fecha=hoy, nota=NOTA_SEMBRADO))
        n += 1
    return n


def notas_v1(notas):
    """Lo que ve el v1: los textos de los hallazgos y las preguntas abiertos, como strings."""
    return {MARCA_V1: TEXTO_MARCA_V1,
            "hallazgos": [h["texto"] for h in notas["hallazgos"] if h.get("estado") == "abierta"],
            "preguntas": [p["texto"] for p in notas["preguntas"] if p.get("estado") == "abierta"]}


def arrastrar_v1(notas, salida):
    """Escribe salidas/<fecha>/notas_tablero.json con las notas abiertas (el v1 lo lee). Si había uno escrito a mano,
    lo guarda una vez como notas_tablero_original.json (sus textos ya pasaron a notas_manuales.json con absorber_v1).
    Devuelve (cantidad de notas arrastradas, si se escribió el archivo)."""
    os.makedirs(salida, exist_ok=True)
    p = os.path.join(salida, NOTAS_V1)
    nuevo = notas_v1(notas)
    if os.path.exists(p):
        viejo = _leer_json(p)
        if isinstance(viejo, dict):
            if not viejo.get(MARCA_V1):
                bk = os.path.join(salida, NOTAS_V1_ORIGINAL)
                if not os.path.exists(bk):
                    shutil.copy2(p, bk)
            for k, v in viejo.items():  # otras claves que tuviera (por ejemplo, metodologia) se conservan
                if k not in ("hallazgos", "preguntas", MARCA_V1):
                    nuevo[k] = v
            if viejo == nuevo:
                return len(nuevo["hallazgos"]) + len(nuevo["preguntas"]), False
    _escribir_json(p, nuevo)
    return len(nuevo["hallazgos"]) + len(nuevo["preguntas"]), True


# ============================================================================ publicación (7.1 pasos 7 y 9)
def estado_publicacion(raiz):
    """salidas/estado_publicacion.json, con los dos tableros y todos sus campos (None si faltan)."""
    p = os.path.join(raiz, ESTADO)
    e = _leer_json(p) if os.path.exists(p) else {}
    for t in RUTAS_FIJAS:
        e.setdefault(t, {})
        for c in CAMPOS_ESTADO:
            e[t].setdefault(c, None)
    return e


def ruta_fija(raiz, tablero):
    return os.path.join(raiz, RUTAS_FIJAS[tablero])


def publicar_fijo(raiz, tablero, origen, meta):
    """Copia el artifact a su ruta fija y actualiza su entrada en estado_publicacion.json (la url no se toca)."""
    destino = ruta_fija(raiz, tablero)
    shutil.copyfile(origen, destino)
    suma = sha256(destino)
    e = estado_publicacion(raiz)
    url = e[tablero].get("url")
    e[tablero].update({k: v for k, v in meta.items() if k != "url"})
    e[tablero].update(url=url, sha256=suma)
    _escribir_json(os.path.join(raiz, ESTADO), e)
    return destino, suma


def guardar_url(raiz, tablero, url):
    url = (url or "").strip()
    if tablero not in RUTAS_FIJAS:
        raise ErrorUsuario(f"Tablero desconocido: {tablero} (tiene que ser v1 o v2)")
    if not re.match(r"^https?://", url):
        raise ErrorUsuario(f"La URL no parece válida: {url!r}")
    e = estado_publicacion(raiz)
    e[tablero]["url"] = url
    _escribir_json(os.path.join(raiz, ESTADO), e)
    return e
